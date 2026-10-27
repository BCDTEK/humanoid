# -*- coding: utf-8 -*-
"""
方案2 数据转换脚本：读取《人形机器人看板数据源.xlsx》4张表，
生成看板 fetch 的 dashboard.json（含 4 表原始数据 + 市场规模聚合 KPI/series）。

由 GitHub Actions 手动触发运行，也可本地运行。
输出: dashboard.json（与 HTML 看板同目录）
"""
import openpyxl, json, os, io, sys
from collections import Counter

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, '人形机器人看板数据源.xlsx')
OUT = os.path.join(HERE, 'dashboard.json')

def num(v):
    try:
        f = float(v)
        return f if f == int(f) else round(f, 4)
    except:
        return None

def clean(v):
    """NA / None / 空 → None（看板显示 NA）"""
    if v is None: return None
    s = str(v).strip()
    if s == '' or s.upper() == 'NA': return None
    return v

wb = openpyxl.load_workbook(SRC, data_only=True)

def rows_of(sheetname):
    ws = wb[sheetname]
    hdr = [c.value for c in ws[1]]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r is None or all(v is None or str(v).strip()=='' for v in r):
            continue
        d = {}
        for i, h in enumerate(hdr):
            if i < len(r):
                d[h] = r[i]
        out.append(d)
    return out

companies = rows_of('① 企业名单')
products  = rows_of('② 产品名单')
financing = rows_of('③ 融资情况')
market    = rows_of('④ 市场规模')

# ===== 市场规模聚合视图：KPI + 多机构趋势 series =====
# 只取全体市场口径，排除单企业/细分实体；出货量统一为万台
def mk_market_series():
    series = {}
    def add(scope, metric, org, y, v, u, f):
        if metric not in ('出货量','市场规模'): return
        key = f'{scope}/{metric}'
        series.setdefault(key, {}).setdefault(org, {'y':[],'v':[],'f':[],'u':u})
        d = series[key][org]
        if y in d['y']: return
        d['y'].append(y); d['v'].append(v); d['f'].append(f)
    for row in market:
        ent = clean(row.get('企业/全体市场'))
        ent = str(ent).strip() if ent else '全体市场'
        if ent != '全体市场': continue
        scope = clean(row.get('范围')) or '全体'
        metric = clean(row.get('指标类型'))
        if metric not in ('出货量','市场规模'): continue
        y = clean(row.get('数据所属年份'))
        if y is None:
            try: y = int(str(y).split('（')[0])
            except: continue
        try: y = int(float(y))
        except: continue
        # 口径识别：备注含"上半年/H1"→ 年份标注为 {year}H1（半年度口径，非整年）
        bm = str(clean(row.get('备注')) or '')
        import re as _re
        if _re.search(r'上半年|H1|上半年度', bm):
            y = f'{y}H1'
        org = clean(row.get('报告机构'))
        v = num(clean(row.get('数值')))
        u = clean(row.get('量纲'))
        f = clean(row.get('预测标记'))
        f = 'F' if f == 'F' else None
        if metric == '出货量' and u == '台' and v is not None:
            v = round(v/10000, 4); u = '万台'
        if v is None: continue
        add(scope, metric, org, y, v, u, f)
    # 排序年份（H1 标注为字符串，统一按字符串排序，保证 2025<2026<2026H1<2030）
    for orgs in series.values():
        for o, s in orgs.items():
            z = sorted(zip([str(x) for x in s['y']], s['v'], s['f']))
            s['y']=[x[0] for x in z]; s['v']=[x[1] for x in z]; s['f']=[x[2] for x in z]
    return series

def pick(series, scope, metric, year, prefs):
    orgs = series.get(f'{scope}/{metric}', {})
    for o in prefs:
        if o in orgs:
            d = orgs[o]
            for i, yy in enumerate(d['y']):
                if str(yy) == str(year) and d['f'][i] != 'F':
                    return d['v'][i], d['u'], o
    for o, d in orgs.items():
        for i, yy in enumerate(d['y']):
            if str(yy) == str(year) and d['f'][i] != 'F':
                return d['v'][i], d['u'], o
    return None, None, None

def pick_pred(series, scope, metric, year, prefs):
    """预测KPI：取 f=='F' 的预测数据"""
    orgs = series.get(f'{scope}/{metric}', {})
    for o in prefs:
        if o in orgs:
            d = orgs[o]
            for i, yy in enumerate(d['y']):
                if str(yy) == str(year) and d['f'][i] == 'F':
                    return d['v'][i], d['u'], o
    for o, d in orgs.items():
        for i, yy in enumerate(d['y']):
            if str(yy) == str(year) and d['f'][i] == 'F':
                return d['v'][i], d['u'], o
    return None, None, None

series = mk_market_series()

def range_kpi(series, scope, metric, year, note):
    """范围型KPI：收集该指标该年份所有预测机构(f=='F')的预测值，
    返回 {t,min,max,unit,orgs,note}；同单位取min~max，跨单位时单列主单位并附备注。"""
    orgs = series.get(f'{scope}/{metric}', {})
    groups = {}  # unit -> list[(val, org)]
    for o, d in orgs.items():
        for i, yy in enumerate(d['y']):
            if str(yy) == str(year) and d['f'][i] == 'F':
                u = d['u']
                groups.setdefault(u, []).append((d['v'][i], o))
    if not groups:
        return {'t': f'{year}{scope}{metric}预测', 'min': None, 'max': None,
                'unit': '', 'orgs': [], 'note': note + '（暂无机构预测数据）'}
    # 主单位：条目最多的单位
    main_u = max(groups, key=lambda u: len(groups[u]))
    vals = [v for v, _ in groups[main_u]]
    orgs_list = [o for _, o in groups[main_u]]
    vmin, vmax = min(vals), max(vals)
    # 其他单位机构作为备注
    others = []
    for u, lst in groups.items():
        if u != main_u:
            others.append(u + '：' + '、'.join(o for _, o in lst))
    extra = ('；另见 ' + '；'.join(others)) if others else ''
    return {'t': f'{year}{scope}{metric}预测', 'min': vmin, 'max': vmax,
            'unit': main_u, 'orgs': orgs_list, 'note': note + extra}

kpis = []
k1 = range_kpi(series, '全球', '出货量', 2026, '全球人形机器人出货量 2026年预测')
k1['t'] = '2026全球出货量预测'
kpis.append(k1)
k2 = range_kpi(series, '中国', '出货量', 2026, '中国人形机器人出货量 2026年预测')
k2['t'] = '2026中国出货量预测'
kpis.append(k2)
k3 = range_kpi(series, '全球', '市场规模', 2026, '全球人形机器人市场规模 2026年预测')
k3['t'] = '2026全球市场规模预测'
kpis.append(k3)
k4 = range_kpi(series, '中国', '市场规模', 2026, '中国人形机器人市场规模 2026年预测')
k4['t'] = '2026中国市场规模预测'
kpis.append(k4)

# ===== 产品竞争格局聚合 =====
def mk_compete():
    prodrows=[]
    for p in products:
        pr = clean(p.get('售价数值')); pu = clean(p.get('售价单位'))
        if pr is not None and pu and str(pu).upper()!='NA':
            prtxt = f'{pr:g}{pu}'
        else:
            prtxt = None
        ship = clean(p.get('出货数值'))
        ship_txt = f'{ship:g}' if isinstance(ship,(int,float)) else None
        spec = clean(p.get('屏幕规格'))
        if spec and len(str(spec))>34: spec = str(spec)[:34]+'…'
        co = clean(p.get('企业名称'))
        if co: co = str(co).split('（')[0]
        prodrows.append({
            'co': co, 'name': clean(p.get('产品名称')),
            'stage': clean(p.get('阶段')), 'price': prtxt,
            'time': clean(p.get('发布时间')), 'ship': ship_txt,
            'screen': clean(p.get('是否带屏')), 'spec': spec
        })
    stage=Counter(); screen=Counter()
    for p in products:
        st = clean(p.get('阶段'))
        st = str(st) if st else '未知'
        stage[st]+=1
        sc = clean(p.get('是否带屏'))
        sc = str(sc) if sc else '未知'
        screen[sc]+=1
    country = Counter()
    for c in companies:
        ct = clean(c.get('国家'))
        if ct: country[str(ct)] += 1
    return {'prodrows': prodrows, 'stage_dist': dict(stage),
            'screen_dist': dict(screen), 'country_dist': dict(country)}

data = {
    'meta': {'title':'人形机器人市场情报看板','badge':'芯视佳专用',
             'updated':'2026-10-21','src_file':'人形机器人看板数据源.xlsx'},
    'companies': companies,
    'products': products,
    'financing': financing,
    'market': market,
    'kpis': kpis,
    'market_series': series,
    'compete': mk_compete(),
}

with open(OUT,'w',encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=1)

print('OK', OUT)
print('companies', len(companies), 'products', len(products), 'financing', len(financing), 'market', len(market))
print('kpis', kpis)
print('series keys', list(series.keys()))
print('compete prodrows', len(data['compete']['prodrows']))
